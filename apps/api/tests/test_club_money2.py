"""Pagamento parziale, mese esente e relatórios della segreteria
(club/fee_splits.py, club/money_reports.py, docs/club-portal.md).

Le promesse: un pagamento parziale divide la quota in una parte pagata
(incasso del mese in cui è entrata, con la sua ricevuta) e un resto che resta
un debito ovunque (quadro del mese, famiglia, avvisi, debiti), e annullarlo
rimette la quota com'era; un mese esente esce dai conti del mese (né emesso,
né debito, né incasso, né mancato incasso) e si annulla con un tocco; la
matrice atleti × mesi e gli altri relatórios tornano con la contabilità; e
la taxa de falta piena riaperta non lascia una quota da 0 € da pagare.
"""

import io
from datetime import date

import pytest
from sqlalchemy import func, select

from app.core.db import utcnow
from app.domains.club import (
    absence,
    accounting,
    fee_splits,
    money_reports,
    month_board,
    months,
    payment_reminders,
)
from app.domains.club import service as club_service
from app.domains.club.models import ClubFee, ClubPaymentProof
from app.domains.club.service import ClubValidationError
from tests.test_club_fees import _athlete, _plans, _user

TODAY = date(2026, 10, 20)


def _this_month() -> accounting.Period:
    """Il mese in cui i pagamenti di questi test sono davvero registrati
    (`paid_at` = adesso)."""
    now = utcnow().date()
    return accounting.parse_period(f"{now.year:04d}-{now.month:02d}", today=now)


async def _monthly_parts(db, athlete_id, period="2026-10") -> list[ClubFee]:
    return list((await db.execute(
        select(ClubFee).where(
            ClubFee.athlete_id == athlete_id, ClubFee.kind == "MONTHLY",
            (ClubFee.period == period) | ClubFee.period.like(f"{period}/r%"),
        ).order_by(ClubFee.period)
    )).scalars())


async def _setup(db, scope):
    admin = await _user(db, scope, "admin")
    await _plans(db, scope, admin)
    parent = await _user(db, scope)
    # 10 anni federali nel 2026: «Ragazzi 8-13», 40 € con scadenza il 10.
    luca = await _athlete(db, scope, parent, birth_date=date(2016, 3, 3), first_name="Luca")
    return admin, parent, luca


async def _partial(db, scope, admin, fee, amount=2500, reason="Paga il resto a fine mese"):
    return await fee_splits.record_partial_payment(
        db, scope=scope, actor_user_id=admin.id, fee_id=fee.id, amount_cents=amount, reason=reason,
        method="CASH", cash_received_by_name="Mister Rui",
    )


async def test_partial_payment_splits_the_fee_and_the_rest_stays_a_debt(db, scope):
    admin, parent, luca = await _setup(db, scope)
    [october] = await _monthly_parts(db, luca.id)
    assert october.amount_cents == 4000

    proof, rest = await _partial(db, scope, admin, october)

    paid, remainder = await _monthly_parts(db, luca.id)
    assert (paid.id, paid.status, paid.amount_cents, paid.payment_method) == (october.id, "PAID", 2500, "CASH")
    assert paid.cash_received_by_name == "Mister Rui" and paid.paid_via_proof_id == proof.id
    assert (remainder.id, remainder.status, remainder.amount_cents, remainder.period) == (rest.id, "DUE", 1500, "2026-10/r1")
    assert remainder.split_from_fee_id == october.id and remainder.split_reason == "Paga il resto a fine mese"
    assert remainder.due_date == october.due_date and remainder.description.endswith("· resto")
    # La ricevuta è della parte pagata.
    assert proof.amount_cents == 2500 and proof.receipt_number

    board = await month_board.club_month(db, scope=scope, period="2026-10", today=TODAY)
    row = next(r for r in board["rows"] if r["athlete_id"] == luca.id)
    assert row["state"] == "PARTIAL"
    assert (row["paid_cents"], row["due_cents"]) == (2500, 1500)
    assert row["fee"]["id"] == remainder.id and row["fee"]["overdue"] is True
    assert [p["id"] for p in row["parts"]] == [october.id, remainder.id]
    t = board["totals"]
    assert (t["expected_cents"], t["collected_cents"], t["missing_cents"]) == (4000, 2500, 1500)
    assert (t["to_pay"], t["partial"], t["overdue"], t["paid"]) == (1, 1, 1, 0)

    # La famiglia vede il resto da pagare, e lo può pagare con una prova.
    family = await month_board.family_month(db, scope=scope, athletes=[luca], today=TODAY)
    assert family["rows"][0]["state"] == "PARTIAL" and family["rows"][0]["fee"]["amount_cents"] == 1500

    # Il calendario della stagione.
    season = await months.season_months(db, athlete=luca, season=2026, today=TODAY)
    cell = next(m for m in season["months"] if m["period"] == "2026-10")
    assert (cell["state"], cell["paid_cents"], cell["due_cents"]) == ("PARTIAL", 2500, 1500)
    assert season["summary"]["paid_cents"] == 2500 and season["summary"]["due_cents"] >= 1500

    # Contabilità: l'incasso è la parte pagata; l'emesso resta la mensilità
    # intera; il resto è un debito.
    month = _this_month()
    figures = accounting._sum(
        await accounting.monthly_figures(db, club_ids=[scope.club_id], start=month.start, end=month.end),
        months=month.months,
    )
    assert figures.income_by_kind.get("MONTHLY") == 2500
    october_figures = accounting._sum(
        await accounting.monthly_figures(db, club_ids=[scope.club_id], start=date(2026, 10, 1), end=date(2026, 11, 1)),
        months=["2026-10"],
    )
    assert october_figures.issued_cents - october_figures.issued_paid_cents >= 1500
    debts = await accounting.debts(db, scope=scope, today=TODAY, limit=None)
    assert any(r["user_id"] == parent.id and r["due_cents"] >= 1500 for r in debts["top"])

    # Gli avvisi della mensilità parlano del resto.
    due = await payment_reminders._due_monthly(db, scope=scope, period="2026-10")
    assert [f.id for f, _a in due] == [remainder.id]

    # Il resto pagato chiude il mese.
    await club_service.staff_record_payment(
        db, scope=scope, actor_user_id=admin.id, fee_ids=[remainder.id], method="MBWAY",
    )
    board = await month_board.club_month(db, scope=scope, period="2026-10", today=TODAY)
    row = next(r for r in board["rows"] if r["athlete_id"] == luca.id)
    assert row["state"] == "PAID" and (row["paid_cents"], row["due_cents"]) == (4000, 0)
    assert board["totals"]["collected_cents"] == 4000


async def test_partial_payment_is_validated(db, scope):
    admin, _parent, luca = await _setup(db, scope)
    [october] = await _monthly_parts(db, luca.id)
    with pytest.raises(ClubValidationError):
        await _partial(db, scope, admin, october, amount=4000)
    with pytest.raises(ClubValidationError):
        await _partial(db, scope, admin, october, amount=0)
    with pytest.raises(ClubValidationError):
        await _partial(db, scope, admin, october, reason="  ")
    # Un contante senza dire a chi: niente resto creato a metà.
    with pytest.raises(ClubValidationError):
        await fee_splits.record_partial_payment(
            db, scope=scope, actor_user_id=admin.id, fee_id=october.id, amount_cents=1000, reason="Resto dopo",
            method="CASH",
        )
    [still] = await _monthly_parts(db, luca.id)
    assert (still.status, still.amount_cents) == ("DUE", 4000)


async def test_undo_partial_payment_puts_the_fee_back_together(db, scope):
    admin, _parent, luca = await _setup(db, scope)
    [october] = await _monthly_parts(db, luca.id)
    proof, _rest = await _partial(db, scope, admin, october)

    fee = await fee_splits.undo_partial_payment(db, scope=scope, fee_id=october.id, actor_user_id=admin.id)

    assert (fee.status, fee.amount_cents, fee.paid_via_proof_id) == ("DUE", 4000, None)
    assert [f.id for f in await _monthly_parts(db, luca.id)] == [october.id]
    assert (await db.get(ClubPaymentProof, proof.id)).status == "REJECTED"
    with pytest.raises(ClubValidationError):
        await fee_splits.undo_partial_payment(db, scope=scope, fee_id=october.id, actor_user_id=admin.id)

    # Anche il rifiuto della ricevuta da «Comprovativos» la ricompone.
    proof2, _rest = await _partial(db, scope, admin, october, amount=1000)
    await club_service.reject_proof(db, scope=scope, proof_id=proof2.id, actor_user_id=admin.id, reason="Errore di cassa")
    [whole] = await _monthly_parts(db, luca.id)
    assert (whole.status, whole.amount_cents) == ("DUE", 4000)


async def test_a_paid_remainder_blocks_the_undo(db, scope):
    admin, _parent, luca = await _setup(db, scope)
    [october] = await _monthly_parts(db, luca.id)
    _proof, rest = await _partial(db, scope, admin, october)
    await club_service.mark_fee_paid(db, scope=scope, fee_id=rest.id, actor_user_id=admin.id, note=None, method="CARD")
    with pytest.raises(ClubValidationError):
        await fee_splits.undo_partial_payment(db, scope=scope, fee_id=october.id, actor_user_id=admin.id)
    # Il resto si può pagare a sua volta in parte: nasce /r2.
    await club_service.reopen_fee(db, scope=scope, fee_id=rest.id, actor_user_id=admin.id)
    _proof, rest2 = await _partial(db, scope, admin, rest, amount=500)
    assert (rest2.period, rest2.amount_cents, rest2.split_from_fee_id) == ("2026-10/r2", 1000, rest.id)


async def test_exempt_month_leaves_the_accounts_of_the_month(db, scope):
    admin, parent, luca = await _setup(db, scope)
    sara = await _athlete(db, scope, parent, birth_date=date(2016, 5, 5), first_name="Sara", sex="F")
    [october] = await _monthly_parts(db, luca.id)

    with pytest.raises(ClubValidationError):
        await fee_splits.exempt_month(db, scope=scope, athlete_id=luca.id, period="2026-10", reason="", actor_user_id=admin.id)
    fee = await fee_splits.exempt_month(
        db, scope=scope, athlete_id=luca.id, period="2026-10", reason="Assente tutto il mese (Erasmus)",
        actor_user_id=admin.id,
    )
    assert (fee.id, fee.status, fee.waive_kind, fee.waive_reason) == (
        october.id, "WAIVED", "EXEMPT", "Assente tutto il mese (Erasmus)",
    )
    with pytest.raises(ClubValidationError):
        await fee_splits.exempt_month(db, scope=scope, athlete_id=luca.id, period="2026-10", reason="Di nuovo",
                                      actor_user_id=admin.id)

    board = await month_board.club_month(db, scope=scope, period="2026-10", today=TODAY)
    rows = {r["first_name"]: r for r in board["rows"]}
    assert rows["Luca"]["state"] == "MONTH_EXEMPT" and rows["Luca"]["fee"]["waive_kind"] == "EXEMPT"
    sara_cents = rows["Sara"]["due_cents"]
    t = board["totals"]
    assert (t["expected_cents"], t["missing_cents"], t["exempt"]) == (sara_cents, sara_cents, 1)

    october_figures = accounting._sum(
        await accounting.monthly_figures(db, club_ids=[scope.club_id], start=date(2026, 10, 1), end=date(2026, 11, 1)),
        months=["2026-10"],
    )
    sara_fee = (await _monthly_parts(db, sara.id))[0]
    issued_without_luca = sum(
        f.amount_cents for f in (await db.execute(select(ClubFee).where(
            ClubFee.club_id == scope.club_id, ClubFee.due_date >= date(2026, 10, 1), ClubFee.due_date < date(2026, 11, 1),
            ClubFee.status != "WAIVED",
        ))).scalars()
    )
    assert october_figures.issued_cents == issued_without_luca
    assert october_figures.exempt_cents == 4000 and october_figures.waived_cents == 0
    assert october_figures.as_dict()["missed_cents"] == october_figures.discounts_cents + october_figures.absence_cents
    debts = await accounting.debts(db, scope=scope, today=TODAY, limit=None)
    assert all(r["due_cents"] != 0 for r in debts["top"])
    due = await payment_reminders._due_monthly(db, scope=scope, period="2026-10")
    assert [f.id for f, _a in due] == [sara_fee.id]

    # Annullare: torna da pagare, con tutto il resto pulito.
    back = await fee_splits.undo_exempt_month(db, scope=scope, fee_id=october.id, actor_user_id=admin.id, today=TODAY)
    assert (back.status, back.waive_kind, back.waive_reason, back.amount_cents) == ("DUE", None, None, 4000)


async def test_exempt_future_month_and_paid_month(db, scope):
    admin, _parent, luca = await _setup(db, scope)
    # Gennaio non è ancora emesso: la mensilità nasce esente, e il job del
    # 1° non ne crea un'altra.
    fee = await fee_splits.exempt_month(
        db, scope=scope, athlete_id=luca.id, period="2027-01", reason="Infortunio", actor_user_id=admin.id,
    )
    assert (fee.period, fee.status, fee.waive_kind) == ("2027-01", "WAIVED", "EXEMPT")
    assert await club_service.generate_monthly_fees(db, scope=scope, period="2027-01", today=TODAY) == 0
    assert (await _monthly_parts(db, luca.id, "2027-01"))[0].status == "WAIVED"
    # Annullare un mese futuro toglie la quota: la emetterà il job.
    assert await fee_splits.undo_exempt_month(db, scope=scope, fee_id=fee.id, actor_user_id=admin.id, today=TODAY) is None
    assert await _monthly_parts(db, luca.id, "2027-01") == []

    [october] = await _monthly_parts(db, luca.id)
    await club_service.mark_fee_paid(db, scope=scope, fee_id=october.id, actor_user_id=admin.id, note=None, method="CARD")
    with pytest.raises(ClubValidationError):
        await fee_splits.exempt_month(db, scope=scope, athlete_id=luca.id, period="2026-10", reason="Assente",
                                      actor_user_id=admin.id)
    # Agosto: il gruppo non fattura, non c'è niente da esentare.
    with pytest.raises(ClubValidationError):
        await fee_splits.exempt_month(db, scope=scope, athlete_id=luca.id, period="2027-08", reason="Vacanze",
                                      actor_user_id=admin.id)


async def test_reopening_a_full_absence_fee_does_not_leave_a_zero_fee(db, scope):
    admin, _parent, luca = await _setup(db, scope)
    [october] = await _monthly_parts(db, luca.id)
    fee = await absence.apply(db, scope=scope, fee_id=october.id, adjustment="FULL", actor_user_id=admin.id)
    assert (fee.status, fee.amount_cents, fee.waived_by_user_id) == ("WAIVED", 0, admin.id)
    figures = accounting._sum(
        await accounting.monthly_figures(db, club_ids=[scope.club_id], start=date(2026, 10, 1), end=date(2026, 11, 1)),
        months=["2026-10"],
    )
    assert figures.absence_cents >= 4000

    fee = await club_service.reopen_fee(db, scope=scope, fee_id=october.id, actor_user_id=admin.id)
    assert (fee.status, fee.amount_cents, fee.absence_adjustment) == ("DUE", 4000, None)


async def test_absence_formula_survives_a_partial_payment(db, scope):
    admin, _parent, luca = await _setup(db, scope)
    [october] = await _monthly_parts(db, luca.id)
    await absence.apply(db, scope=scope, fee_id=october.id, adjustment="HALF", actor_user_id=admin.id)
    before = accounting._sum(
        await accounting.monthly_figures(db, club_ids=[scope.club_id], start=date(2026, 10, 1), end=date(2026, 11, 1)),
        months=["2026-10"],
    )
    await _partial(db, scope, admin, october, amount=1500)  # 20 € di mezza mensilità, 15 pagati
    after = accounting._sum(
        await accounting.monthly_figures(db, club_ids=[scope.club_id], start=date(2026, 10, 1), end=date(2026, 11, 1)),
        months=["2026-10"],
    )
    assert after.absence_cents == before.absence_cents
    assert after.issued_cents == before.issued_cents and after.discounts_cents == before.discounts_cents


async def test_matrix_report_and_its_files(db, scope):
    admin, parent, luca = await _setup(db, scope)
    sara = await _athlete(db, scope, parent, birth_date=date(2016, 5, 5), first_name="Sara", sex="F")
    sara.member_number = "123"
    await db.commit()
    [october] = await _monthly_parts(db, luca.id)
    await _partial(db, scope, admin, october)
    await club_service.generate_monthly_fees(db, scope=scope, period="2026-11", today=TODAY)
    await fee_splits.exempt_month(db, scope=scope, athlete_id=sara.id, period="2026-11", reason="Assente",
                                  actor_user_id=admin.id)

    period = money_reports.report_period(season=2026, year=None, today=TODAY)
    data = await money_reports.build(db, scope=scope, report="matrix", period=period, today=TODAY)
    assert data["months"][0] == "2026-09" and data["months"][-1] == "2027-08" and len(data["months"]) == 12
    rows = {r["first_name"]: r for r in data["rows"]}
    luca_oct = rows["Luca"]["cells"][1]
    assert (luca_oct["state"], luca_oct["paid_cents"], luca_oct["due_cents"], luca_oct["method"]) == ("PARTIAL", 2500, 1500, "CASH")
    assert luca_oct["reason"] == "Paga il resto a fine mese" and luca_oct["receiver"] == "Mister Rui"
    assert rows["Sara"]["cells"][2]["state"] == "MONTH_EXEMPT" and rows["Sara"]["member_number"] == "123"
    assert rows["Luca"]["cells"][0]["state"] == "NOT_ENROLLED"  # iscritto a ottobre
    assert rows["Luca"]["cells"][3]["state"] == "UPCOMING"
    # I totali di riga e di colonna tornano con le caselle.
    for r in data["rows"]:
        assert r["paid_cents"] == sum(c["paid_cents"] for c in r["cells"])
        assert r["due_cents"] == sum(c["due_cents"] for c in r["cells"])
    for i, col in enumerate(data["columns"]):
        assert col["paid_cents"] == sum(r["cells"][i]["paid_cents"] for r in data["rows"])
    assert data["totals"]["paid_cents"] == sum(r["paid_cents"] for r in data["rows"]) == 2500

    csv_text = money_reports.to_csv(data, lang="pt").decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith("Atleta;N.º sócio;Grupo;set 2026")
    assert "Parcial 25,00" in csv_text and "Isento este mês: Assente" in csv_text

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(money_reports.to_xlsx(data, lang="en")))
    sheet = wb.worksheets[0]
    assert sheet.cell(row=2, column=1).value == "Athlete"
    assert len(wb.worksheets) == 2 and wb.worksheets[1].max_row == 1 + 12 * len(data["rows"])


async def test_other_reports_match_the_accounting(db, scope):
    admin, parent, luca = await _setup(db, scope)
    sara = await _athlete(db, scope, parent, birth_date=date(2016, 5, 5), first_name="Sara", sex="F")
    [october] = await _monthly_parts(db, luca.id)
    await _partial(db, scope, admin, october)
    [sara_oct] = await _monthly_parts(db, sara.id)
    await club_service.staff_record_payment(db, scope=scope, actor_user_id=admin.id, fee_ids=[sara_oct.id], method="MBWAY")
    await fee_splits.exempt_month(db, scope=scope, athlete_id=sara.id, period="2026-12", reason="Viaggio",
                                  actor_user_id=admin.id)

    now = utcnow().date()
    period = money_reports.report_period(season=None, year=now.year, today=now)
    month = _this_month()
    figures = accounting._sum(
        await accounting.monthly_figures(db, club_ids=[scope.club_id], start=month.start, end=month.end),
        months=month.months,
    )

    collections = await money_reports.build(db, scope=scope, report="collections", period=period, today=now)
    assert collections["totals"]["total"] == figures.income_cents == 2500 + sara_oct.amount_cents
    assert collections["totals"]["CASH"] == 2500 and collections["totals"]["MBWAY"] == sara_oct.amount_cents

    cash = await money_reports.build(db, scope=scope, report="cash", period=period, today=now)
    assert [(r["receiver"], r["total"]) for r in cash["rows"]] == [("Mister Rui", 2500)]

    debts = await money_reports.build(db, scope=scope, report="debts", period=period, today=TODAY)
    total_due = (await db.execute(select(func.sum(ClubFee.amount_cents)).where(
        ClubFee.club_id == scope.club_id, ClubFee.status == "DUE", ClubFee.amount_cents > 0,
    ))).scalar_one()
    assert debts["totals"]["total"] == total_due
    assert sum(debts["totals"][b] for b, _l, _h in accounting.AGING_BUCKETS) == total_due

    adjustments = await money_reports.build(db, scope=scope, report="adjustments", period=period, today=now)
    types = {r["type"]: r for r in adjustments["rows"]}
    assert types["PARTIAL"]["amount"] == 1500 and types["PARTIAL"]["paid"] == 2500
    assert types["PARTIAL"]["reason"] == "Paga il resto a fine mese"
    assert types["EXEMPT"]["reason"] == "Viaggio" and types["EXEMPT"]["by"]

    pnl = await money_reports.build(db, scope=scope, report="pnl", period=period, today=now)
    assert pnl["totals"]["income"] == figures.income_cents
    assert money_reports.to_csv(pnl, lang="it").decode("utf-8-sig").splitlines()[0].startswith("Mese;Emesso;Incassi")
    assert money_reports.to_xlsx(adjustments, lang="es")[:2] == b"PK"

    with pytest.raises(ClubValidationError):
        money_reports.report_period(season=2026, year=2026, today=now)


async def test_money_routes_through_http_are_for_the_secretary(db, scope, monkeypatch):
    """Le rotte vere: la segreteria (club.manage) incassa in parte, esenta e
    scarica i relatórios; un genitore resta fuori (403)."""
    import redis.asyncio as redis

    from app.core import rate_limit
    from tests.test_club_api import _client

    class _NoRedis:
        async def incr(self, _key):
            raise redis.RedisError("non in questo test")

    # Il limitatore non deve aprire un client Redis legato al loop di questo
    # test (resterebbe nel modulo e romperebbe i test dopo): fallisce aperto.
    monkeypatch.setattr(rate_limit, "_get_redis", lambda: _NoRedis())
    admin, parent, luca = await _setup(db, scope)
    [october] = await _monthly_parts(db, luca.id)

    app, client = await _client(db, monkeypatch, user=admin, scope=scope, granted={"club.manage"})
    try:
        async with client:
            r = await client.post(f"/club/admin/fees/{october.id}/partial-payment", data={
                "amount_cents": "1000", "reason": "Resto la settimana prossima", "method": "MBWAY",
            })
            assert r.status_code == 201, r.text
            assert r.json()["remainder"]["amount_cents"] == 3000
            assert r.json()["remainder"]["split_reason"] == "Resto la settimana prossima"
            r = await client.post("/club/admin/month/exempt", json={
                "athlete_id": str(luca.id), "period": "2026-11", "reason": "Assente",
            })
            assert r.status_code == 200, r.text and r.json()["waive_kind"] == "EXEMPT"
            r = await client.get("/club/admin/reports/matrix?season=2026")
            assert r.status_code == 200 and r.json()["rows"][0]["cells"][1]["state"] == "PARTIAL"
            r = await client.get("/club/admin/reports/collections?season=2026&format=xlsx&lang=pt")
            assert r.status_code == 200 and r.content[:2] == b"PK"
            assert "relatorio-collections-season-2026.xlsx" in r.headers["content-disposition"]
            r = await client.get("/club/admin/reports/debts?format=csv&lang=en")
            assert r.status_code == 200 and r.text.lstrip("\ufeff").startswith("Family,")
            assert (await client.get("/club/admin/reports/nope")).status_code == 404
    finally:
        app.dependency_overrides.clear()

    app, client = await _client(db, monkeypatch, user=parent, scope=scope, granted={"club.family"})
    try:
        async with client:
            assert (await client.get("/club/admin/reports/matrix")).status_code == 403
            assert (await client.post("/club/admin/month/exempt", json={
                "athlete_id": str(luca.id), "period": "2026-12", "reason": "Assente",
            })).status_code == 403
    finally:
        app.dependency_overrides.clear()
